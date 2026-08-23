import io
from odoo import http
from odoo.http import request
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

class ExportController(http.Controller):
    @http.route('/sales_rep/inventory_adjustment/export/<int:adj_id>', type='http', auth='user')
    def export_inventory_adjustment(self, adj_id, **kw):
        adj = request.env['sales.rep.inventory.adjustment'].browse(adj_id)
        if not adj.exists():
            return request.not_found()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Adjustment"

        # Styles
        header_font = Font(bold=True)
        title_font = Font(bold=True, size=14)

        # Title
        ws.merge_cells('A1:E1')
        ws['A1'] = f"Inventory Adjustment: {adj.name}"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')

        # Header Info
        ws['A3'] = "Sales Representative:"
        ws['B3'] = adj.sales_rep_id.name
        ws['A4'] = "Location:"
        ws['B4'] = adj.location_id.display_name
        ws['D3'] = "Date:"
        ws['E3'] = adj.date.strftime('%Y-%m-%d %H:%M:%S') if adj.date else ''

        for cell in ['A3', 'A4', 'D3']:
            ws[cell].font = header_font

        # Table Header
        headers = ['Product', 'On Hand Qty', 'Counted Qty', 'Difference', 'UoM']
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col, value=text)
            cell.font = header_font

        # Lines
        row = 7
        for line in adj.line_ids:
            ws.cell(row=row, column=1, value=line.product_id.display_name)
            ws.cell(row=row, column=2, value=line.theoretical_qty)
            ws.cell(row=row, column=3, value=line.counted_qty)
            ws.cell(row=row, column=4, value=line.difference_qty)
            ws.cell(row=row, column=5, value=line.product_uom_id.name or '')
            row += 1

        # Adjust column widths
        for col_idx, col in enumerate(ws.columns, start=1):
            max_length = 0
            column = get_column_letter(col_idx)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        output = io.BytesIO()
        wb.save(output)
        xlsx_content = output.getvalue()
        output.close()

        filename = f"Inventory_Adjustment_{adj.name}.xlsx"
        return request.make_response(
            xlsx_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename={filename}')
            ]
        )

    @http.route('/sales_rep_management/sales_report_excel/<int:wizard_id>', type='http', auth='user')
    def export_sales_report(self, wizard_id, **kw):
        wizard = request.env['sales.report.wizard'].browse(wizard_id)
        if not wizard.exists():
            return request.not_found()

        orders = wizard.get_orders()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        # Styles
        header_font = Font(bold=True)
        title_font = Font(bold=True, size=14)

        # Show Sales Rep column if multi-select or all
        show_sales_rep = len(wizard.sales_rep_ids) != 1

        # Title
        title_colspan = 7 if show_sales_rep else 6
        ws.merge_cells(f'A1:{get_column_letter(title_colspan)}1')
        ws['A1'] = f"Sales Report: {wizard.date_from} to {wizard.date_to}"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')

        # Table Header
        headers = ['Date']
        if show_sales_rep:
            headers.append('Sales Rep')
        headers.extend(['Sales Order', 'Customer', 'Total before Discount', 'Discount', 'Total after Discount'])
        
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=text)
            cell.font = header_font

        # Data
        row = 4
        for order in orders:
            col = 1
            ws.cell(row=row, column=col, value=order.date_order.strftime('%Y-%m-%d') if order.date_order else ''); col += 1
            if show_sales_rep:
                ws.cell(row=row, column=col, value=order.sales_rep_id.name); col += 1
            ws.cell(row=row, column=col, value=order.name); col += 1
            ws.cell(row=row, column=col, value=order.partner_id.name); col += 1
            ws.cell(row=row, column=col, value=order.amount_total + order.discount_total); col += 1
            ws.cell(row=row, column=col, value=order.discount_total); col += 1
            ws.cell(row=row, column=col, value=order.amount_total); col += 1
            row += 1

        # Totals
        total_label_col = 4 if show_sales_rep else 3
        ws.cell(row=row, column=total_label_col, value="Total").font = header_font
        
        data_start_col = 5 if show_sales_rep else 4
        ws.cell(row=row, column=data_start_col, value=sum(o.amount_total + o.discount_total for o in orders)).font = header_font
        ws.cell(row=row, column=data_start_col + 1, value=sum(o.discount_total for o in orders)).font = header_font
        ws.cell(row=row, column=data_start_col + 2, value=sum(o.amount_total for o in orders)).font = header_font

        # Adjust column widths
        for col_idx, col in enumerate(ws.columns, start=1):
            max_length = 0
            column = get_column_letter(col_idx)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        output = io.BytesIO()
        wb.save(output)
        xlsx_content = output.getvalue()
        output.close()

        filename = f"Sales_Report_{wizard.date_from}_{wizard.date_to}.xlsx"
        return request.make_response(
            xlsx_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename={filename}')
            ]
        )

    @http.route('/sales_rep_management/customer_debt_report_excel/<int:wizard_id>', type='http', auth='user')
    def export_customer_debt_report(self, wizard_id, **kw):
        wizard = request.env['customer.debt.report.wizard'].browse(wizard_id)
        if not wizard.exists():
            return request.not_found()

        partners = wizard.get_debt_data()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customer Debt Report"

        # Styles
        header_font = Font(bold=True)
        title_font = Font(bold=True, size=14)

        # Title
        ws.merge_cells('A1:B1')
        ws['A1'] = "Customer Debt Report"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')

        # Header Info
        ws['A3'] = "Sales Representatives:"
        ws['B3'] = wizard.get_sales_rep_names()
        ws['A4'] = "Date:"
        ws['B4'] = wizard.get_today_date()

        for cell in ['A3', 'A4']:
            ws[cell].font = header_font

        # Table Header
        headers = ['Customer', 'Amount Due']
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col, value=text)
            cell.font = header_font

        # Data
        row = 7
        total_debt = 0.0
        for partner in partners:
            ws.cell(row=row, column=1, value=partner.name)
            ws.cell(row=row, column=2, value=partner._get_total_due())
            total_debt += partner._get_total_due()
            row += 1

        # Totals
        ws.cell(row=row, column=1, value="Total").font = header_font
        ws.cell(row=row, column=2, value=total_debt).font = header_font

        # Adjust column widths
        for col_idx, col in enumerate(ws.columns, start=1):
            max_length = 0
            column = get_column_letter(col_idx)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        output = io.BytesIO()
        wb.save(output)
        xlsx_content = output.getvalue()
        output.close()

        filename = f"Customer_Debt_Report_{wizard.get_today_date()}.xlsx"
        return request.make_response(
            xlsx_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename={filename}')
            ]
        )

    @http.route('/sales_rep_management/collection_report_excel/<int:wizard_id>', type='http', auth='user')
    def export_collection_report(self, wizard_id, **kw):
        wizard = request.env['collection.report.wizard'].browse(wizard_id)
        if not wizard.exists():
            return request.not_found()

        collections = wizard._get_collections()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Collection Report"

        # Styles
        header_font = Font(bold=True)
        title_font = Font(bold=True, size=14)

        # Title
        ws.merge_cells('A1:E1')
        ws['A1'] = "Collection Report"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')

        # Header Info
        ws['A3'] = "Period:"
        ws['B3'] = f"{wizard.date_from} to {wizard.date_to}"
        ws['A4'] = "Sales Representatives:"
        ws['B4'] = wizard.get_sales_rep_names()
        ws['D3'] = "Date:"
        ws['E3'] = wizard.get_today_date()

        for cell in ['A3', 'A4', 'D3']:
            ws[cell].font = header_font

        # Table Header
        headers = ['Date', 'Sales Rep', 'Customer', 'Payment', 'Payment Method']
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col, value=text)
            cell.font = header_font

        # Data
        row = 7
        total_amount = 0.0
        for collection in collections:
            ws.cell(row=row, column=1, value=collection.date.strftime('%Y-%m-%d') if collection.date else '')
            ws.cell(row=row, column=2, value=collection.sales_rep_id.name)
            ws.cell(row=row, column=3, value=collection.partner_id.name)
            ws.cell(row=row, column=4, value=collection.amount)
            ws.cell(row=row, column=5, value=collection.journal_id.name)
            total_amount += collection.amount
            row += 1

        # Totals
        ws.cell(row=row, column=3, value="Total Payment").font = header_font
        ws.cell(row=row, column=4, value=total_amount).font = header_font

        # Adjust column widths
        for col_idx, col in enumerate(ws.columns, start=1):
            max_length = 0
            column = get_column_letter(col_idx)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        output = io.BytesIO()
        wb.save(output)
        xlsx_content = output.getvalue()
        output.close()

        filename = f"Collection_Report_{wizard.date_from}_{wizard.date_to}.xlsx"
        return request.make_response(
            xlsx_content,
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename={filename}')
            ]
        )
